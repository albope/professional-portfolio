from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from theme.tokens import FONTS, GEOMETRY, GRID, PALETTE, dark_accent


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", ""))


def _strip_effects(shape) -> None:
    """Force flat geometry even when the Office theme assigns a shadow style."""
    for effect_ref in shape._element.xpath(".//a:effectRef"):
        effect_ref.set("idx", "0")
    for effect_list in shape._element.xpath(".//a:effectLst"):
        effect_list.getparent().remove(effect_list)


def set_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = None,
    line: str | None = None,
    line_width: float = GEOMETRY.hairline_pt,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    _strip_effects(shape)
    return shape


def line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str,
    width: float = GEOMETRY.hairline_pt,
):
    shape = slide.shapes.add_connector(
        1,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    _strip_effects(shape)
    return shape


def textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    font: str = FONTS.body,
    size: float = 14,
    color: str = PALETTE.ink,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    line_spacing: float | None = None,
    tracking: float | None = None,
    uppercase: bool = False,
    fit: bool = False,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing is not None:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text.upper() if uppercase else text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if tracking is not None:
        # DrawingML stores character spacing in 1/100 pt.
        run._r.get_or_add_rPr().set("spc", str(int(tracking * size * 100)))
    if fit:
        tf.fit_text(font_family=font, max_size=Pt(size))
    return shape


def rich_textbox(
    slide,
    runs: Sequence[dict],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for run_spec in runs:
        run = p.add_run()
        run.text = run_spec["text"]
        run.font.name = run_spec.get("font", FONTS.body)
        run.font.size = Pt(run_spec.get("size", 14))
        run.font.bold = run_spec.get("bold", False)
        run.font.color.rgb = rgb(run_spec.get("color", PALETTE.ink))
    return shape


def paragraphs(
    slide,
    items: Iterable[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    font: str = FONTS.body,
    size: float = 12,
    color: str = PALETTE.ink,
    bullet_color: str | None = None,
    gap: float = 8,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        p.text = f"■  {item}"
        p.font.name = font
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
    return shape


def eyebrow(slide, text: str, *, dark: bool = False, x: float = GRID.left, y: float = 0.54, w: float = 5.0):
    accent = dark_accent(dark)
    rect(slide, x, y + 0.035, 0.075, 0.075, fill=accent)
    return textbox(
        slide,
        text,
        x + 0.17,
        y,
        w,
        0.25,
        font=FONTS.mono,
        size=8.5,
        color=accent,
        tracking=0.18,
        uppercase=True,
    )


def wordmark(slide, *, dark: bool = False, x: float = 10.9, y: float = 0.48, scale: float = 1.0):
    ink = PALETTE.paper if dark else PALETTE.ink
    accent = PALETTE.cobalt_bright if dark else PALETTE.cobalt
    textbox(slide, "BPM", x, y, 0.7 * scale, 0.28 * scale, font=FONTS.display, size=13 * scale, color=ink)
    textbox(slide, "TECH", x + 0.67 * scale, y + 0.065 * scale, 0.55 * scale, 0.2 * scale, font=FONTS.display, size=7.2 * scale, color=accent)
    rect(slide, x + 1.19 * scale, y + 0.14 * scale, 0.08 * scale, 0.08 * scale, fill=ink)


def footer(slide, number: int, total: int, *, dark: bool = False, label: str = "BPM TECH · PROPUESTA"):
    line_color = PALETTE.line_dark if dark else PALETTE.line
    text_color = PALETTE.ink_faint if not dark else "8C8C95"
    line(slide, GRID.left, GRID.footer_y - 0.12, GRID.width - GRID.right, GRID.footer_y - 0.12, color=line_color)
    textbox(slide, label, GRID.left, GRID.footer_y, 4.2, 0.18, font=FONTS.mono, size=7.2, color=text_color, tracking=0.12, uppercase=True)
    textbox(
        slide,
        f"{number:02d} / {total:02d}",
        GRID.width - GRID.right - 1.0,
        GRID.footer_y,
        1.0,
        0.18,
        font=FONTS.mono,
        size=7.2,
        color=text_color,
        align=PP_ALIGN.RIGHT,
    )


def title_block(
    slide,
    title: str,
    *,
    dark: bool = False,
    x: float = GRID.left,
    y: float = 1.02,
    w: float = 8.9,
    h: float = 1.05,
    size: float = 30,
    square_x: float | None = None,
    square_y: float | None = None,
):
    color = PALETTE.paper if dark else PALETTE.ink
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    title_run = p.add_run()
    title_run.text = title.upper()
    title_run.font.name = FONTS.display
    title_run.font.size = Pt(size)
    title_run.font.color.rgb = rgb(color)
    square_run = p.add_run()
    square_run.text = " ■"
    square_run.font.name = "Arial"
    square_run.font.size = Pt(size * 0.30)
    square_run.font.color.rgb = rgb(dark_accent(dark))
    return shape


def add_picture_cover(slide, image_path: str | Path, x: float, y: float, w: float, h: float):
    image_path = Path(image_path)
    with Image.open(image_path) as image:
        img_ratio = image.width / image.height
    frame_ratio = w / h
    if img_ratio > frame_ratio:
        crop_w = frame_ratio / img_ratio
        crop_left = (1 - crop_w) / 2
        crop = (crop_left, 0, crop_left, 0)
    else:
        crop_h = img_ratio / frame_ratio
        crop_top = (1 - crop_h) / 2
        crop = (0, crop_top, 0, crop_top)
    picture = slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), Inches(w), Inches(h))
    picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom = crop
    _strip_effects(picture)
    return picture


def registration_marks(slide, x: float, y: float, w: float, h: float, *, dark: bool = False):
    color = PALETTE.paper if dark else PALETTE.ink
    length = 0.10
    for px, py, dx, dy in (
        (x, y, 1, 1),
        (x + w, y, -1, 1),
        (x, y + h, 1, -1),
        (x + w, y + h, -1, -1),
    ):
        line(slide, px, py, px + dx * length, py, color=color, width=0.6)
        line(slide, px, py, px, py + dy * length, color=color, width=0.6)


def label_value(
    slide,
    label: str,
    value: str,
    x: float,
    y: float,
    w: float,
    *,
    dark: bool = False,
    value_size: float = 13,
):
    accent = dark_accent(dark)
    text_color = PALETTE.paper if dark else PALETTE.ink
    textbox(slide, label, x, y, w, 0.2, font=FONTS.mono, size=7.6, color=accent, tracking=0.15, uppercase=True)
    textbox(slide, value, x, y + 0.28, w, 0.55, font=FONTS.body, size=value_size, color=text_color)
