from pptx import Presentation
from pptx.util import Inches

from components.primitives import footer, set_background, wordmark
from theme.tokens import GRID, PALETTE


def create_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(GRID.width)
    presentation.slide_height = Inches(GRID.height)
    while presentation.slides:
        relationship_id = presentation.slides._sldIdLst[0].rId
        presentation.part.drop_rel(relationship_id)
        del presentation.slides._sldIdLst[0]
    return presentation


def blank_slide(presentation: Presentation, *, dark: bool = False):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    set_background(slide, PALETTE.ink if dark else PALETTE.paper)
    return slide


def chrome(slide, number: int, total: int, *, dark: bool = False, label: str = "BPM TECH · PROPUESTA"):
    wordmark(slide, dark=dark)
    footer(slide, number, total, dark=dark, label=label)
